"""Matrix tests for multi-format ingestion (md, txt, html, docx)."""

import io

from docx import Document as DocxDocument
from packages.core.models import Document, DocumentFormat
from services.ingestion.parsers import parse_file


def _make_docx_bytes(title: str, body: str) -> bytes:
    doc = DocxDocument()
    doc.add_heading(title, level=1)
    doc.add_paragraph(body)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_parse_markdown_file():
    content = b"# Title\n\nSome paragraph text here."
    parsed = parse_file("test.md", content)
    assert parsed.title == "Title"
    assert "Some paragraph" in parsed.text


def test_parse_txt_file():
    content = b"Plain text content here.\n\nSecond paragraph."
    parsed = parse_file("test.txt", content)
    assert parsed is not None
    assert "Plain text content" in parsed.text


def test_parse_html_file():
    content = b"<html><body><h1>Title</h1><p>Content here.</p></body></html>"
    parsed = parse_file("test.html", content)
    assert "Title" in parsed.text
    assert "Content here" in parsed.text
    assert "<html" not in parsed.text


def test_parse_docx_file():
    content = _make_docx_bytes("Doc Title", "Document body content here.")
    parsed = parse_file("test.docx", content)
    assert "Doc Title" in parsed.text
    assert "Document body content" in parsed.text


def test_parse_file_format_detection():
    """Ensure parser correctly detects format from extension."""
    formats = {
        "test.md": DocumentFormat.MARKDOWN,
        "test.txt": DocumentFormat.TXT,
        "test.html": DocumentFormat.HTML,
    }
    for fname, expected_fmt in formats.items():
        content = b"test content"
        parsed = parse_file(fname, content)
        assert parsed.format == expected_fmt, (
            f"{fname}: expected {expected_fmt}, got {parsed.format}"
        )
