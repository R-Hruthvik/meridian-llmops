"""Regression tests for format parsing and PDF quality checks."""

import io

import pytest

from packages.core.errors import ParseError
from services.ingestion.parsers import parse_file
from services.ingestion.parsers.pdf import parse_pdf
from services.ingestion.parsers.txt import decode_text, parse_txt


def test_text_falls_back_to_latin1_with_warning():
    text, warning = decode_text("café".encode("latin-1"))
    assert text == "café"
    assert warning and "latin-1" in warning


def test_txt_splits_paragraphs_on_whitespace_only_blank_lines():
    parsed = parse_txt("first paragraph\n   \nsecond paragraph")
    assert [block.text for block in parsed.blocks] == ["first paragraph", "second paragraph"]


def test_markdown_keeps_fenced_code_lists_and_tables_as_blocks():
    markdown = """# Guide

- first item
- second item

| Name | Value |
| --- | --- |
| retries | 3 |

```python
print('hello')
```
"""
    parsed = parse_file("guide.md", markdown.encode())
    assert [block.kind for block in parsed.blocks] == ["heading", "list", "table", "code"]
    assert parsed.blocks[0].level == 1
    assert "retries" in parsed.blocks[2].text
    assert "print('hello')" in parsed.blocks[3].text


def test_markdown_title_uses_first_h1_then_filename_and_respects_override():
    from_h1 = parse_file("renamed.md", b"## Subsection\n\n# Actual H1\n\nBody")
    filename_fallback = parse_file("my-guide.md", b"## Subsection\n\nBody")
    explicit = parse_file("renamed.md", b"# Actual H1", title="Provided Title")

    assert from_h1.title == "Actual H1"
    assert filename_fallback.title == "My Guide"
    assert explicit.title == "Provided Title"


def test_docx_extracts_heading_and_table_rows():
    from docx import Document

    document = Document()
    document.add_heading("Operations", level=1)
    document.add_paragraph("Restart the service weekly.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Setting"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Retries"
    table.cell(1, 1).text = "3"
    output = io.BytesIO()
    document.save(output)

    parsed = parse_file("operations.docx", output.getvalue())
    assert parsed.title == "Operations"
    assert any(block.kind == "heading" and block.level == 1 for block in parsed.blocks)
    table_blocks = [block for block in parsed.blocks if block.kind == "table"]
    assert len(table_blocks) == 1
    assert "Retries | 3" in table_blocks[0].text


def test_pdf_does_not_drop_headers_seen_on_only_half_the_pages(monkeypatch):
    class Page:
        def __init__(self, text):
            self.text = text

        def extract_text(self):
            return self.text

    class Reader:
        is_encrypted = False

        def __init__(self):
            self.pages = [
                Page("Repeated header\nPage one contains enough readable body text."),
                Page("Repeated header\nPage two contains enough readable body text."),
                Page("Page three contains enough readable body text."),
                Page("Page four contains enough readable body text."),
            ]

    monkeypatch.setattr("pypdf.PdfReader", lambda *_args, **_kwargs: Reader())
    parsed = parse_pdf(b"%PDF-fake", title="Header Test")
    assert any("Repeated header" in block.text for block in parsed.blocks)


def test_encrypted_pdf_without_password_has_clear_parse_error():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt("secret")
    output = io.BytesIO()
    writer.write(output)

    with pytest.raises(ParseError, match="encrypted"):
        parse_file("locked.pdf", output.getvalue())
